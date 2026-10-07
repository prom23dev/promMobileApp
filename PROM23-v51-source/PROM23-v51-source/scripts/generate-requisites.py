"""Generate the downloadable Word sheet from the same approved public page data."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parent.parent
pages = json.loads((root/'data/site-pages.json').read_text())
page = next(page for page in pages['pages'] if page['slug'] == 'details')
doc = Document()
section = doc.sections[0]
section.header.paragraphs[0].text = 'ПРОМЫШЛЕННИК  /  prom23.ru'
section.top_margin = section.bottom_margin = Inches(.65)
style = doc.styles['Normal']; style.font.name = 'Arial'; style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(10)
doc.add_heading('Реквизиты компании', 0)
for paragraph in [p for s in page['sections'] for p in s['paragraphs']]:
    p = doc.add_paragraph(paragraph.replace('Примечание ', ''))
    if paragraph.startswith('Общество'): p.runs[0].bold = True
for style_name in ['Title', 'Heading 1']:
    doc.styles[style_name].font.color.rgb = RGBColor.from_string('202338')
doc.core_properties.title = 'Реквизиты ООО «Промышленник»'
doc.core_properties.author = 'ГК «Промышленник»'
output = root/'mobile/public/documents/prom23-requisites.docx'
output.parent.mkdir(parents=True,exist_ok=True)
doc.save(output)
print(output)
